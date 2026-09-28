from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

A = "Arial"
HDR = Font(name=A, bold=True, color="FFFFFF", size=10)
HDRFILL = PatternFill("solid", fgColor="333333")
BODY = Font(name=A, size=10)
BOLD = Font(name=A, size=10, bold=True)
RED = Font(name=A, size=10, color="C00000", bold=True)
GREY = Font(name=A, size=10, color="808080", italic=True)
TITLE = Font(name=A, size=13, bold=True)
YELL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

BASE = "https://guide.michelin.com/us/en"
FRAME_DATE = "2026-07-28"

# (name, michelin_city_label, tier, url_path, source_page)
P1 = [
 ("Geranium","Copenhagen","3 stars","/capital-region/copenhagen/restaurant/geranium"),
 ("Lille Mølle","Copenhagen","1 star","/capital-region/copenhagen/restaurant/lille-molle"),
 ("formel B","Copenhagen","1 star","/capital-region/copenhagen/restaurant/formel-b"),
 ("Udtryk","Copenhagen","1 star","/capital-region/copenhagen/restaurant/udtryk"),
 ("Uformel","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/uformel"),
 ("Radio","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/radio"),
 ("Norrlyst","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/norrlyst"),
 ("Botschaft","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/botschaft"),
 ("Bobe","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/bobe"),
 ("Gabrielle","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/gabrielle"),
 ("Abigail & Co","Copenhagen","Selected","/capital-region/copenhagen/restaurant/abigail-co"),
 ("Restaurant Glassalen","Copenhagen","Selected","/capital-region/copenhagen/restaurant/restaurant-glassalen"),
 ("Sankt Annæ","Copenhagen","Selected","/capital-region/copenhagen/restaurant/sankt-annae"),
 ("Aamanns 1921","Copenhagen","Selected","/capital-region/copenhagen/restaurant/aamanns-1921"),
 ("Theo","Copenhagen","Selected","/capital-region/copenhagen/restaurant/theo"),
 ("Restaurant Palægade","Copenhagen","Selected","/capital-region/copenhagen/restaurant/palaegade"),
 ("Restaurant Kanalen","Copenhagen","Selected","/capital-region/copenhagen/restaurant/kanalen"),
 ("Mark","Copenhagen","Selected","/capital-region/copenhagen/restaurant/mark"),
 ("Marchal","Copenhagen","1 star","/capital-region/copenhagen/restaurant/marchal"),
 ("Kadeau Copenhagen","Copenhagen","3 stars","/capital-region/copenhagen/restaurant/kadeau-copenhagen"),
 ("Kong Hans Kælder","Copenhagen","2 stars","/capital-region/copenhagen/restaurant/kong-hans-kaelder"),
 ("Alchemist","Copenhagen","2 stars","/capital-region/copenhagen/restaurant/alchemist"),
 ("Koan","Copenhagen","2 stars","/capital-region/copenhagen/restaurant/koan"),
 ("a|o|c","Copenhagen","2 stars","/capital-region/copenhagen/restaurant/a%E2%80%9Ao%E2%80%9Ac"),
 ("akmē","Copenhagen","1 star","/capital-region/copenhagen/restaurant/akme"),
 ("Aure","Copenhagen","1 star","/capital-region/copenhagen/restaurant/restaurant-aure"),
 ("Sushi Anaba","Copenhagen","1 star","/capital-region/copenhagen/restaurant/sushi-anaba"),
 ("ESSE","Copenhagen","1 star","/capital-region/copenhagen/restaurant/esse"),
 ("Alouette","Copenhagen","1 star","/capital-region/copenhagen/restaurant/alouette"),
 ("texture","Copenhagen","1 star","/capital-region/copenhagen/restaurant/texture-1215658"),
 ("Jatak","Copenhagen","1 star","/capital-region/copenhagen/restaurant/jatak"),
 ("Petra Bar & Restaurant","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/petra-bar-restaurant"),
 ("Restaurant Frank","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/frank569608"),
 ("Sanchez","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/sanchez"),
 ("no.2","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/no-2"),
 ("Enomania","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/enomania"),
 ("Silberbauers Bistro","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/silberbauers-bistro"),
 ("Rebel","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/rebel"),
 ("Pluto","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/pluto"),
 ("Anarki","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/anarki"),
 ("Marv & Ben","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/marv-ben"),
 ("Mêlée","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/melee"),
 ("Graziano","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/graziano"),
 ("Selma","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/selma"),
 ("Calma","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/calma"),
 ("Bistro Lupa","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/bistro-lupa"),
 ("Koefoed","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/koefoed"),
 ("Paesàno","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/paesano"),
]

P2 = [
 ("Levi","Copenhagen","Bib Gourmand","/capital-region/copenhagen/restaurant/levi"),
 ("Fasangården","Copenhagen","Selected","/capital-region/copenhagen/restaurant/fasangarden"),
 ("Restaurant Anton","Copenhagen","Selected","/capital-region/copenhagen/restaurant/restaurant-anton"),
 ("Ark","Copenhagen","Selected","/capital-region/copenhagen/restaurant/ark"),
 ("Kødbyens Fiskebar","Copenhagen","Selected","/capital-region/copenhagen/restaurant/kodbyens-fiskebar"),
 ("Høst","Copenhagen","Selected","/capital-region/copenhagen/restaurant/host"),
 ("Barr","Copenhagen","Selected","/capital-region/copenhagen/restaurant/barr"),
 ("The Pescatarian","Copenhagen","Selected","/capital-region/copenhagen/restaurant/the-pescatarian"),
 ("Kiin Kiin VeVe","Copenhagen","Selected","/capital-region/copenhagen/restaurant/kiin-kiin-veve"),
 ("Aotori","Copenhagen","Selected","/capital-region/copenhagen/restaurant/kappo-ando"),
 ("Kiin Kiin","Copenhagen","Selected","/capital-region/copenhagen/restaurant/kiin-kiin"),
 ("Møntergade","Copenhagen","Selected","/capital-region/copenhagen/restaurant/montergade"),
 ("Krogs Fiskerestaurant","Copenhagen","Selected","/capital-region/copenhagen/restaurant/krogs-fiskerestaurant"),
 ("Magny","Copenhagen","Selected","/capital-region/copenhagen/restaurant/magny"),
 ("à terre","Copenhagen","Selected","/capital-region/copenhagen/restaurant/a-terre"),
 ("Amalie","Copenhagen","Selected","/capital-region/copenhagen/restaurant/amalie"),
 ("Vækst","Copenhagen","Selected","/capital-region/copenhagen/restaurant/vaekst"),
 ("Admiralgade 26","Copenhagen","Selected","/capital-region/copenhagen/restaurant/admiralgade-26"),
 ("Mielcke & Hurtigkarl","Copenhagen","Selected","/capital-region/copenhagen/restaurant/mielcke-hurtigkarl"),
 ("Restaurant VIE","Nordhavn","Bib Gourmand","/capital-region/nordhavn_7770461/restaurant/restaurant-vie"),
 ("Jordnær","Gentofte","3 stars","/capital-region/gentofte_1373269/restaurant/jordnaer"),
 ("Parsley Salon","Hellerup","1 star","/capital-region/hellerup_7770408/restaurant/parsley-salon"),
 ("The Samuel","Hellerup","1 star","/capital-region/hellerup_7770408/restaurant/the-samuel"),
 ("Søllerød Kro","Holte","1 star","/capital-region/holte_1374367/restaurant/sollerod-kro"),
 ("Vollmers","Malmo","2 stars","/skane/malmo/restaurant/vollmers"),
 ("Västergatan","Malmo","Bib Gourmand","/skane/malmo/restaurant/vastergatan"),
 ("Ruths","Malmo","Bib Gourmand","/skane/malmo/restaurant/ruths"),
 ("Namu","Malmo","Bib Gourmand","/skane/malmo/restaurant/namu"),
 ("Mutantur","Malmo","Bib Gourmand","/skane/malmo/restaurant/mutantur"),
 ("Kockeriet","Malmo","Selected","/skane/malmo/restaurant/kockeriet"),
 ("Bloom in the Park","Malmo","Selected","/skane/malmo/restaurant/bloom-in-the-park"),
 ("Lyran Matbar","Malmo","Selected","/skane/malmo/restaurant/lyran"),
 ("aster","Malmo","Selected","/skane/malmo/restaurant/aster-1203662"),
 ("Restaurang Atmosfär","Malmo","Selected","/skane/malmo/restaurant/atmosfar"),
 ("Brasserie Sture 1912","Malmo","Selected","/skane/malmo/restaurant/brasserie-sture"),
]

# F1 verdict from the guide's own city label. PROVISIONAL — F0/F1 proper needs address resolution.
OUT_LABELS = {"Gentofte": "Gentofte Kommune", "Hellerup": "Gentofte Kommune",
              "Holte": "Rudersdal Kommune", "Malmo": "Malmö, Sweden"}

ROWS = []
for lst, page in ((P1, 1), (P2, 2)):
    for (n, city, tier, path) in lst:
        out = OUT_LABELS.get(city)
        ROWS.append(dict(name=n, city=city, tier=tier, url=BASE + path, page=page,
                         excl="outside_boundary" if out else "",
                         kommune=out if out else "TO RESOLVE"))

assert len(ROWS) == 83, len(ROWS)
tiers = {}
for r in ROWS:
    tiers[r["tier"]] = tiers.get(r["tier"], 0) + 1
STATED = {"3 stars": 3, "2 stars": 5, "1 star": 14, "Bib Gourmand": 29, "Selected": 32}
assert tiers == STATED, (tiers, STATED)   # D-032 count assertion
print("count assertion PASSED:", tiers, "total", len(ROWS))

COLS = [("venue_name",24),("michelin_city_label",19),("route_a1",9),("route_a1_tier",14),
        ("kommune",20),("address",34),("google_place_id",30),("neighbourhood",14),
        ("route_a2",9),("route_a2_tier",13),("route_b",9),("route_b_citations",16),
        ("ownership_scale",15),("ownership_count",15),("ownership_source",18),("parent_group",15),
        ("opened_on",11),("menu_publication_mode",21),("menu_source_url",30),
        ("cheapest_full_meal_price",23),("currency",9),
        ("in_cohort",10),("excluded_reason",18),
        ("frame_date",11),("source_url",62),("source_page",12),("observed_at",12),
        ("retrieved_at",12),("added_by",10),("notes",46)]

wb = Workbook()
ws = wb.active; ws.title = "cph_route_a1_frame"

ws["A1"] = "COPENHAGEN — ROUTE A1 FRAME (MICHELIN Guide Nordic Countries 2026)"
ws["A1"].font = TITLE
ws["A2"] = ("Route A1 only. Route A2 (White Guide) and Route B not yet applied — blank means not looked at, never 'no'. "
            "F1 verdicts below are PROVISIONAL, taken from the guide's own city label, not from a resolved address.")
ws["A2"].font = GREY

r = 4
for i,(nm,w) in enumerate(COLS, start=1):
    c = ws.cell(row=r, column=i, value=nm); c.font = HDR; c.fill = HDRFILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[r].height = 30
ws.freeze_panes = "B5"

r = 5
for d in ROWS:
    obs = "2026-07-28" if d["page"] == 1 else "2026-07-30"
    note = ""
    if d["excl"]:
        note = "Fails F1 on the guide's city label."
        if d["city"] == "Malmo":
            note += " Sweden — Michelin's 'Copenhagen and surroundings' crosses the Øresund."
    elif d["city"] == "Nordhavn":
        note = "Nordhavn is inside Københavns Kommune. Guide files it as a separate locality."
    if d["name"] == "Marchal":
        note = "venue_type = hotel_restaurant (Hotel d'Angleterre). Qualifies on A1 in its own right per D-028."
    if d["name"] == "formel B":
        note = "Address is 1800 Frederiksberg C — Frederiksberg Kommune, in scope. Guide labels it 'Copenhagen'."
    vals = [d["name"], d["city"], "TRUE", d["tier"], d["kommune"], "TO RESOLVE", "TO RESOLVE", "",
            "", "", "", "", "", "", "", "", "", "", "", "", "",
            "FALSE", d["excl"], FRAME_DATE, d["url"], f"page {d['page']}", obs, "2026-07-30", "claude", note]
    for i,v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v); c.font = BODY; c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=(i in (25,30)))
    if d["excl"]:
        ws.cell(row=r, column=23).font = RED
        ws.cell(row=r, column=5).font = RED
    r += 1
last = r - 1

r += 1
ws.cell(row=r, column=1, value="COUNTS — every one states its denominator").font = BOLD; r += 1
S = [("A1 candidates enumerated (guide's stated total: 83)", f'=COUNTA(A5:A{last})'),
     ("  3 stars (guide states 3)", f'=COUNTIF(D5:D{last},"3 stars")'),
     ("  2 stars (guide states 5)", f'=COUNTIF(D5:D{last},"2 stars")'),
     ("  1 star (guide states 14)", f'=COUNTIF(D5:D{last},"1 star")'),
     ("  Bib Gourmand (guide states 29)", f'=COUNTIF(D5:D{last},"Bib Gourmand")'),
     ("  Selected (guide states 32)", f'=COUNTIF(D5:D{last},"Selected")'),
     ("Provisionally OUT on city label — outside_boundary", f'=COUNTIF(W5:W{last},"outside_boundary")'),
     ("  of which Malmö, Sweden", f'=COUNTIF(B5:B{last},"Malmo")'),
     ("Provisionally IN, pending address resolution", f'=COUNTA(A5:A{last})-COUNTIF(W5:W{last},"outside_boundary")')]
for lab,f in S:
    ws.cell(row=r, column=1, value=lab).font = BODY
    c = ws.cell(row=r, column=2, value=f); c.font = BOLD; c.fill = YELL
    r += 1
r += 1
ws.cell(row=r, column=1, value=("D-032 count assertion PASSED: all five tier counts and the total reconcile "
                                "against the guide's own stated figures.")).font = BOLD

# legend
lg = wb.create_sheet("legend_and_provenance")
lg.column_dimensions["A"].width = 30; lg.column_dimensions["B"].width = 100
lg["A1"] = "How to read this file"; lg["A1"].font = TITLE
rows = [
 ("What this is","The complete Route A1 pool for Copenhagen from MICHELIN Guide Nordic Countries 2026, enumerated in full."),
 ("What it is not","Not the frame. A2 and Route B are not applied, so venues qualifying only on those routes are absent."),
 ("Source","guide.michelin.com, Copenhagen listing, pages 1 and 2. Retrieved via Firecrawl, proxy=basic, HTTP 200. "
           "No stealth mode, no proxy rotation — robots.txt grants named AI crawlers unrestricted access (06 §4)."),
 ("Count assertion (D-032)","The guide states 83 restaurants and gives per-tier facet counts of 3/5/14/29/32. "
           "Rows parsed: 83, tiers 3/5/14/29/32. Exact match on all six figures. The run would have halted on any mismatch."),
 ("Parsing (D-031)","Pages scraped to markdown and parsed deterministically. No LLM extraction was used to decide what is on the list."),
 ("observed_at differs by page","Page 1 came from Firecrawl's cache dated 2026-07-28 — the frame date exactly. Page 2 was fetched "
           "live on 2026-07-30 with maxAge=0. Both dates sit inside the same guide edition (published 1 June 2026), so the "
           "selection is identical; the dates are recorded per row rather than smoothed."),
 ("F1 is PROVISIONAL","Verdicts come from the guide's own city label, not a resolved address. Doc 21 §3 requires F0 resolution "
           "to a street address and google_place_id before F1 is applied. The label is known to be unreliable: Michelin's own "
           "ceremony article calls Søllerød Kro 'Copenhagen' though it is in Rudersdal Kommune, 15km away."),
 ("Two label traps already found","formel B is labelled 'Copenhagen' but sits in Frederiksberg Kommune — in scope, and would be "
           "wrongly kept for the wrong reason. Nordhavn is filed as a separate locality but is inside Københavns Kommune — "
           "Restaurant VIE is in scope and a naive city-label filter would drop it."),
 ("Blank cells","Not looked at. Never 'no'."),
 ("in_cohort","FALSE on every row. No cohort has been drawn; the seed 20260728 has not been used."),
]
r = 2
for a,b in rows:
    lg.cell(row=r, column=1, value=a).font = BOLD
    c = lg.cell(row=r, column=2, value=b); c.font = BODY
    c.alignment = Alignment(wrap_text=True, vertical="top"); lg.row_dimensions[r].height = 42
    r += 1

wb.save("/mnt/user-data/outputs/cph-route-a1-frame.xlsx")
print("saved")
