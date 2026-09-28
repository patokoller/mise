# Session f00a0b22 turn 13 build_fixture.py, with the two turn-17 str_replace edits applied (akmē resolved; operator verification row).
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
HDR = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
HDRFILL = PatternFill("solid", fgColor="333333")
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
WARN = Font(name=ARIAL, size=10, color="C00000", bold=True)
GREY = Font(name=ARIAL, size=10, color="808080", italic=True)
TITLE = Font(name=ARIAL, size=13, bold=True)
YELL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

FRAME_DATE = "2026-07-28"
RETRIEVED = "2026-07-30"

COLS = [
    ("venue_name", 22), ("address", 42), ("kommune", 18), ("neighbourhood", 14),
    ("google_place_id", 30), ("route_a1", 9), ("route_a1_tier", 13),
    ("route_a2", 9), ("route_a2_tier", 13), ("route_b", 9), ("route_b_citations", 16),
    ("ownership_scale", 15), ("ownership_count", 15), ("ownership_source", 22),
    ("parent_group", 16), ("opened_on", 11),
    ("menu_publication_mode", 21), ("menu_source_url", 46),
    ("cheapest_full_meal_price", 23), ("currency", 9),
    ("in_cohort", 10), ("excluded_reason", 18),
    ("frame_date", 11), ("added_by", 10), ("retrieved_at", 12), ("notes", 60),
]

U = "untested"
N = ""

# venue, address, kommune, place_id, a1_tier, mode, price, menu_url, ownership, own_count, own_src, parent, excl, notes
ROWS = [
 ("akmē","Sandkaj 39, 2150 København","Københavns Kommune","ChIJ97znST9TUkYRXKSLIQ5ezzc","1 star",
  "price_only",1500,"https://www.akme.dk/menu",U,N,N,N,N,
  "RESOLVED BY OPERATOR 2026-07-30: page carried two price blocks, 1500 and 1300 DKK, with companion wine lists dated 29-07-2026 and 05-11-2025. Operator confirmed 1500 is live; 1300 is stale markup still in the DOM. No mechanical rule currently decides this case."),
 ("Alchemist","Refshalevej 173C, 1432 København","Københavns Kommune","ChIJGbFWVvBSUkYR5x0i6Z836Ao","2 stars",
  "price_only",None,N,U,N,N,N,N,
  "NOT RETESTED 2026-07-30. Mode carried from the calibration record in 21 §11.1 / D-027. Price never recorded there. Re-test before use."),
 ("Alouette","Kronprinsessegade 8, 1306 København","Københavns Kommune","ChIJcXrr90dTUkYRNJo6gjI26uk","1 star",
  "price_only",2195,"https://www.restaurantalouette.dk/menu",U,N,N,N,N,
  "Three hops: root is a splash page, then /home, then /menu. Menu page carries meta-ROBOTS: NOINDEX — posture question for doc 10."),
 ("AOC","Dronningens Tværgade 2, 1302 København","Københavns Kommune","ChIJM6fyhhhTUkYRy5Y4y78NEYU","2 stars",
  "full",2300,"https://restaurantaoc.dk/en/menu",U,N,N,N,N,
  "Two menus: AOC TASTING MENU 3600, THE MENU 2300. Dish lists run together on the page; the shorter menu repeats a subset of the longer one's dishes."),
 ("Aure","Krudtløbsvej 8, 1439 København","Københavns Kommune","ChIJV_8E0sJTUkYRRlJfxBk5d6k","1 star",
  "full",1850,"https://www.restaurantaure.dk/menu",U,N,N,N,N,
  "Clean case. Full dish list, single price, one hop from homepage."),
 ("ESSE","Trelleborggade 13a, 2100 København","Københavns Kommune","ChIJZ7egZwBTUkYRAQ5WQBRR2nM","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested — outside the sample."),
 ("formel B","Vesterbrogade 182, 1800 Frederiksberg","Frederiksberg Kommune","ChIJtxaAuZlTUkYRt9S0ULtIiuA","1 star",
  "full",1500,"https://formelbgroup.dk/formelb/en/food/","small_group",6,"formelbgroup.dk, read 2026-07-30","formel B Group",N,
  "Three aggregators gave three different official domains (formelfamily.dk, formel-b.dk, formelbgroup.dk). Canonical site resolution is a step Places does not provide. Group operates 6 venues per its own footer."),
 ("Geranium","Per Henrik Lings Allé 4, 8., 2100 København","Københavns Kommune","ChIJAQshwflSUkYRF9f4wpDIt7U","3 stars",
  "price_only",4400,"https://www.geranium.dk/en/menu",U,N,N,N,N,
  "Single-page site: /en/menu returns the homepage. Publishes menu name and price only. Page also still carries a COVID-era notice, so freshness cannot be inferred from the page."),
 ("JATAK","Rantzausgade 39, 2200 København","Københavns Kommune","ChIJu7H8HXxTUkYRuLKQbQU824w","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested — outside the sample."),
 ("Kadeau","Wildersgade 10B, 1408 København","Københavns Kommune","ChIJ1Y52MZlTUkYRuDl4FNhYh_U","3 stars",
  U,None,N,U,N,N,N,N,"F2 not tested. Promoted to 3 stars in the 2026 edition."),
 ("Koan","Langeliniekaj 5, 2100 København","Københavns Kommune","ChIJoxbQhIBTUkYR15NWp69nCt0","2 stars",
  U,None,N,U,N,N,N,N,"F2 not tested — outside the sample."),
 ("Kong Hans Kælder","Vingårdstræde 6, 1070 København","Københavns Kommune","ChIJXdbkpBdTUkYRO6NpiMExk-w","2 stars",
  U,None,N,U,N,N,N,N,"F2 not tested. Places returns this venue under the English name 'King Hans Cellar' — a name-matching case for F0."),
 ("Lille Mølle","Christianshavns Voldgade 52, 1424 København","Københavns Kommune","ChIJ29LGogVTUkYRcswK32qPzAY","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested. New star in the 2026 edition."),
 ("Marchal","Hotel d'Angleterre, Kongens Nytorv 34, 1050 København","Københavns Kommune","ChIJtRm9KRhTUkYRhIT68DTO0r8","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested. venue_type = hotel_restaurant; qualifies on A1 in its own right per F6 as amended (D-028)."),
 ("Sushi Anaba","Mariehamngade 23, 2150 København","Københavns Kommune","ChIJc7_JfGlTUkYRbVK5FOovyOs","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested — outside the sample."),
 ("Texture","Sølvgade 86, 1307 København","Københavns Kommune","ChIJRxVkAFJTUkYRGaQE2cWK7fU","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested. Michelin styles this venue lowercase as 'texture'."),
 ("Udtryk","Teglgårdstræde 8A, 1452 København","Københavns Kommune","ChIJn1w7Q6ZTUkYRHVs90CYuMPg","1 star",
  U,None,N,U,N,N,N,N,"F2 not tested — outside the sample."),
 ("Jordnær","Gentoftegade 29, 2820 Gentofte","Gentofte Kommune","ChIJWxyiJPFNUkYRTp10-YnYcNg","3 stars",
  N,None,N,U,N,N,N,"outside_boundary",
  "Fails F1. Gentofte Kommune. Michelin's own 2026 list files this under Denmark; Danish press treats it as Copenhagen."),
 ("Parsley Salon","Strandvejen 203, 2900 Hellerup","Gentofte Kommune","ChIJhTbZQNNTUkYRrlURja_vdgM","1 star",
  N,None,N,U,N,N,N,"outside_boundary","Fails F1. Gentofte Kommune. Michelin lists it as 'Hellerup'."),
 ("The Samuel","Hellerupvej 40, 2900 Hellerup","Gentofte Kommune","ChIJ97nnHBpTUkYRtJSEQUisX6o","1 star",
  N,None,N,U,N,N,N,"outside_boundary","Fails F1. Gentofte Kommune. Michelin lists it as 'Copenhagen'."),
 ("Søllerød Kro","Søllerødvej 35, 2840 Søllerød","Rudersdal Kommune","ChIJua9HeudOUkYRDn953dEtd1Y","1 star",
  N,None,N,U,N,N,N,"outside_boundary",
  "Fails F1. Rudersdal Kommune, ~15km north. Michelin's own list labels it 'Copenhagen'. The clearest demonstration that the guide's city label and the kommune boundary are different objects."),
]

wb = Workbook()

ws = wb.active
ws.title = "f2_test_fixture"

ws["A1"] = "F2 MECHANICAL-DETERMINATION TEST — NOT FRAME DATA"
ws["A1"].font = TITLE
ws["A2"] = ("Test fixture only. Source list is secondary press, not the MICHELIN Guide. "
            "These rows must never be merged into the frame. Purpose: measure whether F2 can be "
            "determined without a human, and how often it cannot.")
ws["A2"].font = GREY

r = 4
for i, (name, width) in enumerate(COLS, start=1):
    c = ws.cell(row=r, column=i, value=name)
    c.font = HDR; c.fill = HDRFILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[r].height = 30
ws.freeze_panes = "A5"

r = 5
for (name, addr, kom, pid, tier, mode, price, murl, oscale, ocount, osrc, parent, excl, notes) in ROWS:
    vals = [
        name, addr, kom, "not recorded", pid,
        "TRUE", tier, U, "", U, "",
        oscale, ocount, osrc, parent, "not recorded",
        mode, murl, price, "DKK" if price else "",
        "FALSE", excl, FRAME_DATE, "claude", RETRIEVED, notes,
    ]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = BODY
        c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=(i in (2, 18, 26)))
    if excl:
        ws.cell(row=r, column=22).font = WARN
    if mode == U:
        ws.cell(row=r, column=17).font = GREY
    if "AMBIGUOUS" in notes or "NOT RETESTED" in notes:
        ws.cell(row=r, column=26).font = WARN
    r += 1

last = r - 1

r += 1
ws.cell(row=r, column=1, value="COUNTS — every one states its denominator").font = BOLD
r += 1
summary = [
    ("Candidates enumerated (starred tier, secondary source)", f'=COUNTA(A5:A{last})'),
    ("Pass F1 (Københavns + Frederiksberg Kommune)", f'=COUNTIF(V5:V{last},"")'),
    ("Fail F1 — outside_boundary", f'=COUNTIF(V5:V{last},"outside_boundary")'),
    ("F2 tested this session, of those passing F1", f'=COUNTIF(Q5:Q{last},"full")+COUNTIF(Q5:Q{last},"price_only")'),
    ("  of which  full", f'=COUNTIF(Q5:Q{last},"full")'),
    ("  of which  price_only", f'=COUNTIF(Q5:Q{last},"price_only")'),
    ("  of which  none", f'=COUNTIF(Q5:Q{last},"none")'),
    ("F2 not tested", f'=COUNTIF(Q5:Q{last},"untested")'),
]
for label, f in summary:
    ws.cell(row=r, column=1, value=label).font = BODY
    c = ws.cell(row=r, column=2, value=f)
    c.font = BOLD; c.fill = YELL
    r += 1

r += 1
ws.cell(row=r, column=1, value=(
    "Note: one of the price_only rows (Alchemist) was carried from the 21 §11.1 calibration record and "
    "not retested on 2026-07-30. Six venues were tested fresh this session.")).font = WARN

lg = wb.create_sheet("legend")
lg.column_dimensions["A"].width = 30
lg.column_dimensions["B"].width = 95
lg["A1"] = "How to read this file"; lg["A1"].font = TITLE
rows = [
    ("", ""),
    ("What this is", "A test of whether filter F2 can be decided by machine. It is not the frame and not a cohort."),
    ("Source of the venue list", "The 21 Copenhagen starred restaurants in MICHELIN Guide Nordic Countries 2026, as reported by "
                                 "Michelin's own ceremony article and two independent outlets. A complete tier was used so that no "
                                 "venue was hand-picked."),
    ("Why a complete tier", "Choosing venues by hand would put the operator's intuition inside the measurement device, which is "
                            "the failure doc 21 §1 exists to prevent."),
    ("What is verified", "F0 (entity resolution), F1 (kommune boundary), address and google_place_id — all 21 rows, via Google Places."),
    ("What is partly done", "F2 — six venues tested fresh, one carried from the calibration record, ten untested."),
    ("What is absent", "Route A2, Route B, ownership for 20 of 21, opened_on, neighbourhood, cohort membership. "
                       "Blank means not looked at, never 'no'."),
    ("in_cohort", "FALSE on every row. No cohort has been drawn and the seed has not been used."),
    ("Prices", "Cheapest price at which one person can eat a complete meal as the venue structures it, per doc 21 §5. "
               "Recorded as a figure, never a band."),
    ("Red text", "A row a human has to look at. These are the finding, not a defect in the file."),
    ("Grey text", "Not tested."),
]
r = 2
for a, b in rows:
    lg.cell(row=r, column=1, value=a).font = BOLD
    c = lg.cell(row=r, column=2, value=b); c.font = BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    lg.row_dimensions[r].height = 30
    r += 1

pv = wb.create_sheet("provenance")
pv.column_dimensions["A"].width = 26
pv.column_dimensions["B"].width = 95
pv["A1"] = "Provenance and dates"; pv["A1"].font = TITLE
prov = [
    ("frame_date", FRAME_DATE + " — frozen per doc 21 §2. Not changed by this test."),
    ("retrieved_at", RETRIEVED + " — every fact in this file was read on this date."),
    ("observed_at", "Not recorded. Venue pages do not date their menus, which is itself a finding: "
                    "observed_at cannot be derived from a menu page and must come from the retrieval event."),
    ("Venue list source", "Secondary press reporting of MICHELIN Guide Nordic Countries 2026 (ceremony 1 June 2026, "
                          "Tivoli Concert Hall). guide.michelin.com itself was unreachable — see 06 §4, reach = waf_blocked."),
    ("Cross-check", "21 restaurants / 31 stars agrees across Michelin's own ceremony article, VisitCopenhagen and "
                    "two independent outlets, and the star arithmetic reconciles (3x3 + 4x2 + 14x1 = 31)."),
    ("Operator verification", "2026-07-30, by eye, against the live pages: Geranium price_only CONFIRMED, "
                              "AOC full CONFIRMED, akmē price 1500 DKK CONFIRMED. Two of two classifications the "
                              "operator could check independently were correct; the third row was a case the rule "
                              "could not decide and the operator decided it."),
    ("Entity resolution", "Google Places, one query per venue, 2026-07-30. All 21 resolved to exactly one trading venue; "
                          "no F0 unresolved rows."),
    ("F1 method", "Postcode and kommune from the resolved Places address, against doc 21 §3 F1 "
                  "(Københavns Kommune + Frederiksberg Kommune)."),
    ("F4", "NOT tested. Google Places did not return business_status in these responses, so 'trading at the frame date' "
           "rests on the venue having published hours. This needs a separate source."),
    ("What Places did not give", "No website URL and no business_status. Every venue site had to be found by a separate "
                                 "search, which is an extra step and an extra failure point in any automated enumerator."),
]
r = 3
for a, b in prov:
    pv.cell(row=r, column=1, value=a).font = BOLD
    c = pv.cell(row=r, column=2, value=b); c.font = BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    pv.row_dimensions[r].height = 30
    r += 1

wb.save("/mnt/user-data/outputs/f2-test-fixture-copenhagen.xlsx")
print("saved; data rows:", len(ROWS))
